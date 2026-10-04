import os
import pymupdf

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext in (".txt", ".md"):
        return _extract_plain(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".pptx":
        return _extract_pptx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_path):
    doc = pymupdf.open(file_path)
    pages = []
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text()
        if text.strip():
            pages.append({"text": text, "page": page_num})
    return pages


def _extract_plain(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    return [{"text": text, "page": 1}]


def _extract_docx(file_path):
    from docx import Document
    doc = Document(file_path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [{"text": text, "page": 1}]


def _extract_pptx(file_path):
    from pptx import Presentation
    prs = Presentation(file_path)
    pages = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        text_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        text_parts.append(line)
        text = "\n".join(text_parts)
        if text.strip():
            pages.append({"text": text, "page": slide_num})
    return pages